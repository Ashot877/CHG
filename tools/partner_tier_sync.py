import re
import unicodedata
from collections import defaultdict
from email import policy
from email.parser import BytesParser
from io import BytesIO
from html.parser import HTMLParser
import xml.etree.ElementTree as ET
from urllib.parse import parse_qs, urlparse

import pandas as pd
from docx import Document
import requests
import streamlit as st
from requests.auth import HTTPBasicAuth

from core.config import PARTNER_TIER_CONFIG
from core.jira import (
    clean_base_url,
    jira_get_field_ids_by_name,
    jira_get_issue_editmeta,
    jira_search,
    jira_update_issue_fields,
    summarize_exception,
)
from core.session import current_jira_context, require_jira_settings
from core.utils import compact, field_to_text


DEFAULT_PARTNER_TIER_JQL = '''project = "Change Management"
AND issuetype in (Change)
AND "Partner tier" is EMPTY
AND status NOT IN ("Internal Review", "On Hold", Draft)
AND "Partner / Project" NOT IN (Digitain, "Account Management")
AND created > 2025-01-01
AND (resolution NOT IN ("Wrong ticket", "Not Actual") OR resolution IS EMPTY)
AND "Product Change Category (Partner)" != "Partner Request - Onboarding"
AND approvers IS NOT EMPTY
AND Approvers NOT IN inactiveUsers()
AND status WAS Open'''


class _HTMLTableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self._table_depth = 0
        self._rows = None
        self._row = None
        self._cell = None
        self._cell_is_header = False
        self._cell_rowspan = 1
        self._cell_colspan = 1

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "table":
            self._table_depth += 1
            if self._table_depth == 1:
                self._rows = []
        elif self._table_depth == 1 and tag == "tr":
            self._row = []
        elif self._table_depth == 1 and tag in ("td", "th"):
            self._cell = []
            self._cell_is_header = tag == "th"
            attrs = dict(attrs or [])
            try:
                self._cell_rowspan = max(1, int(attrs.get("rowspan", 1) or 1))
            except (TypeError, ValueError):
                self._cell_rowspan = 1
            try:
                self._cell_colspan = max(1, int(attrs.get("colspan", 1) or 1))
            except (TypeError, ValueError):
                self._cell_colspan = 1
        elif self._table_depth == 1 and tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_data(self, data):
        if self._table_depth == 1 and self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self._table_depth == 1 and tag in ("td", "th") and self._cell is not None:
            text = " ".join("".join(self._cell).replace("\xa0", " ").split())
            if self._row is not None:
                self._row.append({
                    "text": text,
                    "header": self._cell_is_header,
                    "rowspan": self._cell_rowspan,
                    "colspan": self._cell_colspan,
                })
            self._cell = None
            self._cell_rowspan = 1
            self._cell_colspan = 1
        elif self._table_depth == 1 and tag == "tr":
            if self._row is not None and any(cell.get("text") for cell in self._row):
                self._rows.append(self._row)
            self._row = None
        elif tag == "table":
            if self._table_depth == 1 and self._rows:
                self.tables.append(self._rows)
            self._table_depth = max(0, self._table_depth - 1)
            if self._table_depth == 0:
                self._rows = None


def normalize_space(value):
    value = unicodedata.normalize("NFKC", str(value or "")).replace("\xa0", " ")
    return " ".join(value.split()).strip()


def normalize_header(value):
    value = normalize_space(value).casefold()
    value = re.sub(r"[^\w]+", " ", value, flags=re.UNICODE)
    return " ".join(value.split())


def normalize_partner(value):
    value = normalize_space(value).casefold()
    value = re.sub(r"\s*([/\\|_-])\s*", r"\1", value)
    return value


def normalize_partner_loose(value):
    value = unicodedata.normalize("NFKC", str(value or "")).casefold()
    return "".join(ch for ch in value if ch.isalnum())


def extract_tier_from_category(category):
    text = normalize_space(category)
    if not text:
        return None, "Partner Category is empty"

    group_match = re.search(r"\bgroup\s*[-:#]?\s*(\d{1,2})\b", text, flags=re.IGNORECASE)
    if group_match:
        return int(group_match.group(1)), ""

    numbers = re.findall(r"(?<!\d)(\d{1,2})(?!\d)", text)
    unique_numbers = list(dict.fromkeys(numbers))
    if len(unique_numbers) == 1:
        return int(unique_numbers[0]), ""
    if not unique_numbers:
        return None, f'No tier number found in Partner Category: "{text}"'
    return None, f'Ambiguous Partner Category contains several numbers: "{text}"'


def _expand_html_table(rows):
    """Expand HTML rowspan/colspan cells into a rectangular logical table.

    Confluence Word exports commonly use a vertically merged Partner cell.
    Without expanding rowspan, project rows slide one column to the left and
    the parser starts treating a project as a partner. Corporate HTML doing
    corporate HTML things, basically.
    """
    expanded = []
    active_spans = {}

    for raw_row in rows:
        logical = {}
        next_active = {}

        # Carry values from cells that started on a previous row.
        for col, (remaining, cell) in active_spans.items():
            if remaining <= 0:
                continue
            logical[col] = dict(cell)
            if remaining - 1 > 0:
                next_active[col] = (remaining - 1, cell)

        col = 0
        for raw_cell in raw_row:
            while col in logical:
                col += 1

            cell = {
                "text": normalize_space(raw_cell.get("text")),
                "header": bool(raw_cell.get("header")),
            }
            rowspan = max(1, int(raw_cell.get("rowspan", 1) or 1))
            colspan = max(1, int(raw_cell.get("colspan", 1) or 1))

            for offset in range(colspan):
                target_col = col + offset
                while target_col in logical:
                    target_col += 1
                logical[target_col] = dict(cell)
                if rowspan > 1:
                    next_active[target_col] = (rowspan - 1, cell)
            col += colspan

        if logical:
            max_col = max(logical)
            expanded.append([
                logical.get(i, {"text": "", "header": False})
                for i in range(max_col + 1)
            ])
        active_spans = next_active

    return expanded


def _parse_html_tables(html):
    parser = _HTMLTableParser()
    parser.feed(html or "")
    return [_expand_html_table(table) for table in parser.tables]


def _row_texts(row):
    return [normalize_space(cell.get("text")) for cell in row]


def _find_named_columns(rows, partner_header, category_header, project_header="Project name"):
    wanted_partner = normalize_header(partner_header)
    wanted_category = normalize_header(category_header)
    wanted_project = normalize_header(project_header)
    for row_index, row in enumerate(rows[:15]):
        headers = [normalize_header(cell.get("text")) for cell in row]
        partner_indexes = [i for i, h in enumerate(headers) if h == wanted_partner]
        category_indexes = [i for i, h in enumerate(headers) if h == wanted_category]
        if partner_indexes and category_indexes:
            project_indexes = [i for i, h in enumerate(headers) if h == wanted_project]
            return row_index, partner_indexes[0], category_indexes[0], (project_indexes[0] if project_indexes else None)
    return None


def _find_group_column_heuristic(rows):
    if len(rows) < 2:
        return None
    max_cols = max(len(row) for row in rows)
    candidates = []
    for col in range(max_cols):
        values = []
        for row in rows:
            if col < len(row):
                values.append(normalize_space(row[col].get("text")))
        group_hits = sum(bool(re.search(r"\bgroup\s*[-:#]?\s*\d{1,2}\b", v, re.IGNORECASE)) for v in values if v)
        if group_hits:
            candidates.append((group_hits, col))
    if not candidates:
        return None
    candidates.sort(reverse=True)
    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
        return None
    return candidates[0][1]


def parse_partner_rows_from_html(
    html,
    partner_header="Partner",
    category_header="Partner Category",
    project_header="Project name",
):
    tables = _parse_html_tables(html)
    best_rows = []
    best_score = -1

    for rows in tables:
        named = _find_named_columns(rows, partner_header, category_header, project_header)
        parsed = []
        score = 0
        if named:
            header_row, partner_col, category_col, project_col = named
            current_partner = ""
            for row in rows[header_row + 1:]:
                texts = _row_texts(row)
                partner = texts[partner_col] if partner_col < len(texts) else ""
                if partner:
                    current_partner = partner
                else:
                    partner = current_partner
                project = texts[project_col] if project_col is not None and project_col < len(texts) else ""
                category = texts[category_col] if category_col < len(texts) else ""
                if partner or project or category:
                    parsed.append({
                        "Partner": partner,
                        "Project name": project,
                        "Partner Category": category,
                    })
            score = 1000 + len(parsed)
        else:
            category_col = _find_group_column_heuristic(rows)
            if category_col is not None and category_col != 0:
                for row in rows:
                    texts = _row_texts(row)
                    partner = texts[0] if texts else ""
                    category = texts[category_col] if category_col < len(texts) else ""
                    if partner and extract_tier_from_category(category)[0] is not None:
                        parsed.append({"Partner": partner, "Project name": "", "Partner Category": category})
                score = len(parsed)

        if parsed and score > best_score:
            best_rows = parsed
            best_score = score

    if not best_rows:
        raise ValueError(
            f'Could not find a Confluence table with "{partner_header}" and "{category_header}" columns.'
        )
    return best_rows


def extract_page_id(page_url):
    parsed = urlparse((page_url or "").strip())
    query = parse_qs(parsed.query)
    if query.get("pageId"):
        return query["pageId"][0]
    match = re.search(r"/pages/(\d+)(?:/|$)", parsed.path)
    if match:
        return match.group(1)
    return ""


def derive_confluence_base_url(page_url):
    parsed = urlparse((page_url or "").strip())
    if not parsed.scheme or not parsed.netloc:
        return ""
    path = parsed.path or ""
    if path == "/wiki" or path.startswith("/wiki/"):
        return f"{parsed.scheme}://{parsed.netloc}/wiki"
    return f"{parsed.scheme}://{parsed.netloc}"


def _confluence_request_auth(auth_type, username, token):
    auth_type = (auth_type or "").strip().lower()
    headers = {"Accept": "application/json, text/html;q=0.9,*/*;q=0.8"}
    auth = None
    if not token:
        return headers, auth
    if auth_type == "basic":
        auth = HTTPBasicAuth(username or "", token)
    elif auth_type == "bearer":
        headers["Authorization"] = f"Bearer {token}"
    else:
        headers["Authorization"] = f"Bearer {token}"
    return headers, auth


def load_confluence_partner_rows(page_url, jira_auth=None):
    page_url = (page_url or "").strip()
    if not page_url:
        raise ValueError("Confluence page URL is empty.")

    own_token = str(PARTNER_TIER_CONFIG.get("confluence_token", "") or "").strip()
    own_username = str(PARTNER_TIER_CONFIG.get("confluence_username", "") or "").strip()
    own_auth_type = str(PARTNER_TIER_CONFIG.get("confluence_auth_type", "Bearer") or "Bearer").strip()
    inherit = bool(PARTNER_TIER_CONFIG.get("inherit_jira_credentials", False))

    if own_token:
        auth_type, username, token = own_auth_type, own_username, own_token
    elif inherit and jira_auth:
        auth_type, username, token = jira_auth
    else:
        auth_type, username, token = "", "", ""

    headers, auth = _confluence_request_auth(auth_type, username, token)
    page_id = extract_page_id(page_url)
    base_url = clean_base_url(PARTNER_TIER_CONFIG.get("confluence_base_url") or derive_confluence_base_url(page_url))
    errors = []

    if page_id and base_url:
        endpoints = [
            (f"{base_url}/rest/api/content/{page_id}", {"expand": "body.view,body.storage,version"}),
            (f"{base_url}/api/v2/pages/{page_id}", {"body-format": "storage"}),
        ]
        for endpoint, params in endpoints:
            try:
                response = requests.get(endpoint, headers=headers, auth=auth, params=params, timeout=45)
                if not response.ok:
                    errors.append(f"{response.status_code} from {endpoint}")
                    continue
                payload = response.json()
                title = payload.get("title") or "Confluence page"
                html = ""
                body = payload.get("body") or {}
                if isinstance(body, dict):
                    view = body.get("view") or {}
                    storage = body.get("storage") or {}
                    if isinstance(view, dict):
                        html = view.get("value") or ""
                    if not html and isinstance(storage, dict):
                        html = storage.get("value") or ""
                if html:
                    rows = parse_partner_rows_from_html(
                        html,
                        PARTNER_TIER_CONFIG.get("partner_column_name", "Partner"),
                        PARTNER_TIER_CONFIG.get("category_column_name", "Partner Category"),
                        PARTNER_TIER_CONFIG.get("project_column_name", "Project name"),
                    )
                    return rows, title
            except Exception as error:
                errors.append(str(error))

    try:
        response = requests.get(page_url, headers=headers, auth=auth, timeout=45)
        if response.ok:
            rows = parse_partner_rows_from_html(
                response.text,
                PARTNER_TIER_CONFIG.get("partner_column_name", "Partner"),
                PARTNER_TIER_CONFIG.get("category_column_name", "Partner Category"),
                PARTNER_TIER_CONFIG.get("project_column_name", "Project name"),
            )
            return rows, "Confluence page"
        errors.append(f"{response.status_code} from page URL")
    except Exception as error:
        errors.append(str(error))

    detail = "; ".join(errors[-3:]) or "Unknown Confluence error"
    raise RuntimeError(
        "Could not read the Confluence page. Configure [partner_tier] Confluence credentials in secrets.toml "
        f"or enable inherit_jira_credentials. Details: {detail}"
    )


def _rows_from_matrix(
    matrix,
    partner_header="Partner",
    category_header="Partner Category",
    project_header="Project name",
):
    """Extract Partner/Project/Partner Category rows from a rectangular text matrix."""
    rows = [
        [{"text": normalize_space(cell), "header": False} for cell in row]
        for row in matrix
        if any(normalize_space(cell) for cell in row)
    ]
    if not rows:
        return []

    named = _find_named_columns(rows, partner_header, category_header, project_header)
    if not named:
        return []

    header_row, partner_col, category_col, project_col = named
    parsed = []
    current_partner = ""
    for row in rows[header_row + 1:]:
        texts = _row_texts(row)
        partner = texts[partner_col] if partner_col < len(texts) else ""
        if partner:
            current_partner = partner
        else:
            partner = current_partner
        project = texts[project_col] if project_col is not None and project_col < len(texts) else ""
        category = texts[category_col] if category_col < len(texts) else ""
        if partner or project or category:
            parsed.append({
                "Partner": partner,
                "Project name": project,
                "Partner Category": category,
            })
    return parsed


def parse_partner_rows_from_docx(
    data,
    partner_header="Partner",
    category_header="Partner Category",
    project_header="Project name",
):
    document = Document(BytesIO(data))
    best_rows = []
    for table in document.tables:
        matrix = [[cell.text for cell in row.cells] for row in table.rows]
        parsed = _rows_from_matrix(matrix, partner_header, category_header, project_header)
        if len(parsed) > len(best_rows):
            best_rows = parsed

    if not best_rows:
        raise ValueError(
            f'Could not find a Word table with "{partner_header}" and "{category_header}" columns.'
        )
    return best_rows


DOCX_MAGIC = b"PK\x03\x04"
OLE_DOC_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_docx_bytes(data):
    return bool(data) and data.startswith(DOCX_MAGIC)


def _decode_legacy_word_text(data):
    """Decode Word-compatible HTML/XML .doc exports."""
    if not data:
        return ""

    prefix = data[:8192].decode("latin-1", errors="ignore")
    charset_match = re.search(r"charset\s*=\s*[\"']?([a-zA-Z0-9._-]+)", prefix, re.IGNORECASE)
    encodings = []
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.append("utf-16")
    if charset_match:
        encodings.append(charset_match.group(1))
    encodings.extend(["utf-8-sig", "utf-8", "cp1252", "latin-1"])

    seen = set()
    for encoding in encodings:
        key = str(encoding).lower()
        if key in seen:
            continue
        seen.add(key)
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1", errors="ignore")


def _extract_mhtml_html(data):
    try:
        message = BytesParser(policy=policy.default).parsebytes(data)
    except Exception:
        return ""

    parts = message.walk() if message.is_multipart() else [message]
    for part in parts:
        try:
            if str(part.get_content_type()).lower() != "text/html":
                continue
            content = part.get_content()
            if isinstance(content, bytes):
                return _decode_legacy_word_text(content)
            return str(content or "")
        except Exception:
            continue
    return ""


def _html_matrices(html):
    matrices = []
    for table in _parse_html_tables(html):
        matrix = [_row_texts(row) for row in table]
        if matrix:
            matrices.append(matrix)
    return matrices


def _xml_local_name(tag):
    return str(tag or "").rsplit("}", 1)[-1].split(":")[-1].lower()


def _word_xml_matrices(text):
    try:
        root = ET.fromstring(text)
    except Exception:
        return []

    matrices = []
    for table in root.iter():
        if _xml_local_name(table.tag) != "tbl":
            continue
        matrix = []
        for row in table.iter():
            if _xml_local_name(row.tag) != "tr":
                continue
            cells = []
            for cell in list(row):
                if _xml_local_name(cell.tag) != "tc":
                    continue
                cells.append(normalize_space(" ".join(cell.itertext())))
            if cells:
                matrix.append(cells)
        if matrix:
            matrices.append(matrix)
    return matrices


def extract_legacy_word_table_matrices(data):
    """Extract tables from Confluence .doc exports without guessing Jira values.

    Confluence often gives an HTML/MHTML file with a .doc extension. Word XML is
    also supported. A true binary Word 97-2003 file is rejected explicitly.
    """
    if is_docx_bytes(data):
        raise ValueError("DOCX payload should be parsed with the .docx parser.")

    if data.startswith(OLE_DOC_MAGIC):
        raise ValueError(
            "This is a true binary Word 97-2003 .doc file. Confluence exports are usually HTML-based and work directly. "
            "For this binary file, open it in Word and Save As .docx before upload."
        )

    mhtml_html = _extract_mhtml_html(data)
    if mhtml_html:
        matrices = _html_matrices(mhtml_html)
        if matrices:
            return matrices

    text = _decode_legacy_word_text(data)
    if re.search(r"<\s*(?:html|table|!doctype)\b", text, re.IGNORECASE):
        matrices = _html_matrices(text)
        if matrices:
            return matrices

    matrices = _word_xml_matrices(text)
    if matrices:
        return matrices

    raise ValueError(
        "Could not read tables from this .doc file. The helper supports Confluence Word-compatible .doc exports, DOCX, Excel, and CSV."
    )


def parse_partner_rows_from_doc(
    data,
    partner_header="Partner",
    category_header="Partner Category",
    project_header="Project name",
):
    # Some exports use a .doc name for OOXML bytes. Content detection is safer than trusting the extension.
    if is_docx_bytes(data):
        return parse_partner_rows_from_docx(data, partner_header, category_header, project_header)

    best_rows = []
    for matrix in extract_legacy_word_table_matrices(data):
        parsed = _rows_from_matrix(matrix, partner_header, category_header, project_header)
        if len(parsed) > len(best_rows):
            best_rows = parsed

    if not best_rows:
        raise ValueError(
            f'Could not find a Word table with "{partner_header}" and "{category_header}" columns.'
        )
    return best_rows


def _dataframe_matrix(df):
    return [["" if pd.isna(value) else str(value) for value in row] for row in df.values.tolist()]


def parse_partner_rows_from_excel(
    data,
    partner_header="Partner",
    category_header="Partner Category",
    project_header="Project name",
):
    book = pd.ExcelFile(BytesIO(data))
    best_rows = []
    best_sheet = ""
    for sheet_name in book.sheet_names:
        df = pd.read_excel(book, sheet_name=sheet_name, header=None, dtype=object)
        parsed = _rows_from_matrix(_dataframe_matrix(df), partner_header, category_header, project_header)
        if len(parsed) > len(best_rows):
            best_rows = parsed
            best_sheet = str(sheet_name)

    if not best_rows:
        raise ValueError(
            f'Could not find an Excel sheet with "{partner_header}" and "{category_header}" columns.'
        )
    return best_rows, best_sheet


def parse_partner_rows_from_csv(
    data,
    partner_header="Partner",
    category_header="Partner Category",
    project_header="Project name",
):
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            df = pd.read_csv(BytesIO(data), header=None, dtype=object, sep=None, engine="python", encoding=encoding)
            parsed = _rows_from_matrix(_dataframe_matrix(df), partner_header, category_header, project_header)
            if parsed:
                return parsed
        except Exception as error:
            last_error = error
    raise ValueError(
        f'Could not find CSV columns "{partner_header}" and "{category_header}".'
        + (f" Details: {last_error}" if last_error else "")
    )


def load_partner_rows_from_upload(uploaded_file):
    if uploaded_file is None:
        raise ValueError("Source file is not uploaded.")

    name = str(getattr(uploaded_file, "name", "") or "source").strip()
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    data = uploaded_file.getvalue()
    partner_header = str(PARTNER_TIER_CONFIG.get("partner_column_name", "Partner") or "Partner")
    project_header = str(PARTNER_TIER_CONFIG.get("project_column_name", "Project name") or "Project name")
    category_header = str(PARTNER_TIER_CONFIG.get("category_column_name", "Partner Category") or "Partner Category")

    if suffix == "doc":
        rows = parse_partner_rows_from_doc(data, partner_header, category_header, project_header)
        return rows, f"{name} · Confluence Word export"
    if suffix == "docx":
        rows = parse_partner_rows_from_docx(data, partner_header, category_header, project_header)
        return rows, f"{name} · Word export"
    if suffix in ("xlsx", "xlsm"):
        rows, sheet = parse_partner_rows_from_excel(data, partner_header, category_header, project_header)
        title = f"{name} · sheet: {sheet}" if sheet else name
        return rows, title
    if suffix == "csv":
        rows = parse_partner_rows_from_csv(data, partner_header, category_header, project_header)
        return rows, name

    raise ValueError("Unsupported source file. Upload .doc, .docx, .xlsx, .xlsm, or .csv.")


def _collapse_entries(entries):
    valid = [entry for entry in entries if entry.get("tier") is not None]
    if not valid:
        return None, "missing_category"
    tiers = {entry["tier"] for entry in valid}
    if len(tiers) > 1:
        return None, "conflict"
    return valid[0], "ok"


def partner_match_candidates(value):
    """Return safe lookup candidates from Jira's combined Partner / Project value.

    Jira often renders the field as ``Partner - Project``. A common case is
    ``Festwin - Festwin`` while the source table contains only ``Festwin`` in
    Partner and ``-`` in Project name. That repeated pair is safe to collapse.

    Examples:
      GrandPashaBet-ALL -> GrandPashaBet
      albatross-Solibet -> Solibet
      Festwin - Festwin -> Festwin

    We never use fuzzy substring matching. Every derived candidate still has to
    resolve to one unambiguous Partner or Project row in the source file.
    """
    raw = normalize_space(value)
    if not raw:
        return []

    candidates = []
    seen = set()
    # Include the ordinary hyphen plus common Unicode dashes used by copied Jira
    # and Confluence text. Humans have apparently invented several characters
    # for the same tiny horizontal line, so the matcher has to care.
    separator_re = r"\s*[-‐‑‒–—_/|]\s*"
    strip_chars = "-‐‑‒–—_/| "

    def add(candidate, origin):
        candidate = normalize_space(candidate).strip(strip_chars)
        if not candidate:
            return
        dedupe = normalize_partner_loose(candidate)
        if not dedupe or dedupe in seen:
            return
        seen.add(dedupe)
        candidates.append((candidate, origin))

    add(raw, "Original")

    # Jira can display one logical name twice, e.g. ``Festwin - Festwin``.
    # Collapse only when both sides normalize to exactly the same value. This is
    # deterministic and cannot turn GrandPashaBet - Solibet into GrandPashaBet.
    pair_parts = [normalize_space(part) for part in re.split(separator_re, raw) if normalize_space(part)]
    if len(pair_parts) == 2 and normalize_partner_loose(pair_parts[0]) == normalize_partner_loose(pair_parts[1]):
        add(pair_parts[0], "Repeated pair → base")

    all_match = re.match(
        rf"^(.*?)(?:{separator_re}|\s+)ALL\s*$",
        raw,
        flags=re.IGNORECASE,
    )
    base = raw
    if all_match and normalize_space(all_match.group(1)):
        base = normalize_space(all_match.group(1))
        add(base, "ALL → base")

    # A Jira value can carry a family/prefix before the actual project, e.g.
    # albatross-Solibet. Try progressively shorter suffixes, but only if they
    # later resolve to a unique Partner or Project source row.
    parts = [normalize_space(part) for part in re.split(separator_re, base) if normalize_space(part)]
    if len(parts) > 1:
        for index in range(1, len(parts)):
            add("-".join(parts[index:]), "Suffix")

    return candidates


def build_partner_lookup(rows):
    partner_exact = defaultdict(list)
    partner_loose = defaultdict(list)
    project_exact = defaultdict(list)
    project_loose = defaultdict(list)
    prepared = []

    for row in rows:
        partner = normalize_space(row.get("Partner"))
        project = normalize_space(row.get("Project name"))
        category = normalize_space(row.get("Partner Category"))
        if not partner and not project:
            continue
        tier, category_error = extract_tier_from_category(category)
        entry = {
            "partner": partner,
            "project": project,
            "category": category,
            "tier": tier,
            "category_error": category_error,
        }
        prepared.append(entry)
        if partner:
            partner_exact[normalize_partner(partner)].append(entry)
            partner_loose[normalize_partner_loose(partner)].append(entry)
        if project and project not in {"-", "—", "–"}:
            project_exact[normalize_partner(project)].append(entry)
            project_loose[normalize_partner_loose(project)].append(entry)

    return {
        "partner_exact": partner_exact,
        "partner_loose": partner_loose,
        "project_exact": project_exact,
        "project_loose": project_loose,
        "rows": prepared,
    }


def _resolve_tier_match(entries, source_kind, match_name):
    if not entries:
        return None, "", match_name

    if source_kind == "Project":
        parent_partners = {normalize_partner(entry.get("partner")) for entry in entries if entry.get("partner")}
        if len(parent_partners) > 1:
            return None, "Project name match is ambiguous across multiple partners", match_name
    else:
        source_partners = {normalize_partner(entry.get("partner")) for entry in entries if entry.get("partner")}
        if len(source_partners) > 1:
            return None, "Partner match is ambiguous in the source file", match_name

    entry, state = _collapse_entries(entries)
    if state == "conflict":
        return None, f"Conflicting Partner Category values for matched {source_kind.lower()}", match_name
    if state == "missing_category":
        first = entries[0]
        return None, first.get("category_error") or "Partner Category is empty", match_name
    return entry, "", match_name


def match_partner(lookup, partner_value):
    partner_value = normalize_space(partner_value)
    if not partner_value:
        return None, "Partner / Project is empty", ""

    found_errors = []
    for candidate, origin in partner_match_candidates(partner_value):
        exact_key = normalize_partner(candidate)
        loose_key = normalize_partner_loose(candidate)
        origin_suffix = "" if origin == "Original" else f" · {origin}"

        checks = [
            (lookup["project_exact"].get(exact_key, []), "Project", f"Project exact{origin_suffix}"),
            (lookup["partner_exact"].get(exact_key, []), "Partner", f"Partner exact{origin_suffix}"),
            (lookup["project_loose"].get(loose_key, []), "Project", f"Project normalized{origin_suffix}"),
            (lookup["partner_loose"].get(loose_key, []), "Partner", f"Partner normalized{origin_suffix}"),
        ]

        for entries, source_kind, match_name in checks:
            if not entries:
                continue
            entry, error, resolved_match = _resolve_tier_match(entries, source_kind, match_name)
            if not error:
                return entry, "", resolved_match
            found_errors.append((error, resolved_match))

    if found_errors:
        return None, found_errors[0][0], found_errors[0][1]
    return None, "Partner / Project not found in Partner or Project name columns", ""


def _field_values(value):
    if value is None:
        return []
    if isinstance(value, list):
        result = []
        for item in value:
            text = field_to_text(item).strip()
            if text:
                result.append(text)
        return result
    text = field_to_text(value).strip()
    return [text] if text else []


def _is_empty_field(value):
    return value is None or value == "" or value == [] or value == {}


def build_preview_rows(issues, base_url, partner_field_id, tier_field_id, lookup):
    actionable = []
    blocked = []
    for issue in issues:
        fields = issue.get("fields", {}) or {}
        key = issue.get("key", "")
        common = {
            "Key": key,
            "Open": f"{clean_base_url(base_url)}/browse/{key}",
            "Summary": compact(fields.get("summary", ""), 160),
            "Status": field_to_text(fields.get("status")),
        }

        current_tier = fields.get(tier_field_id)
        if not _is_empty_field(current_tier):
            blocked.append({**common, "Partner": field_to_text(fields.get(partner_field_id)), "Reason": "Partner tier is already filled"})
            continue

        partner_values = _field_values(fields.get(partner_field_id))
        if not partner_values:
            blocked.append({**common, "Partner": "", "Reason": "Partner / Project is empty"})
            continue
        if len(partner_values) != 1:
            blocked.append({**common, "Partner": ", ".join(partner_values), "Reason": "Partner / Project has multiple values; automatic choice is disabled"})
            continue

        partner_value = partner_values[0]
        entry, error, match_type = match_partner(lookup, partner_value)
        if error:
            blocked.append({**common, "Partner": partner_value, "Reason": error})
            continue

        actionable.append({
            "Select": True,
            **common,
            "Partner": partner_value,
            "Source Partner": entry["partner"],
            "Source Project": entry.get("project", ""),
            "Partner Category": entry["category"],
            "New tier": int(entry["tier"]),
            "Match": match_type,
        })
    return actionable, blocked


def _allowed_value_payload(field_meta, tier):
    tier_text = str(int(tier))
    allowed = field_meta.get("allowedValues") or [] if isinstance(field_meta, dict) else []
    for item in allowed:
        if not isinstance(item, dict):
            continue
        raw = str(item.get("value") or item.get("name") or "").strip()
        raw_tier, _ = extract_tier_from_category(raw)
        if raw == tier_text or raw_tier == int(tier):
            if item.get("id") is not None:
                return {"id": str(item.get("id"))}
            if item.get("value") is not None:
                return {"value": item.get("value")}
    return None


def candidate_tier_payloads(editmeta, field_id, tier):
    field_meta = (((editmeta or {}).get("fields") or {}).get(field_id) or {}) if isinstance(editmeta, dict) else {}
    allowed_payload = _allowed_value_payload(field_meta, tier)
    candidates = []
    if allowed_payload is not None:
        candidates.append(allowed_payload)

    schema = field_meta.get("schema") or {}
    schema_type = str(schema.get("type") or "").lower()
    custom = str(schema.get("custom") or "").lower()
    if schema_type in ("number", "integer"):
        candidates.append(int(tier))
    elif schema_type == "string":
        candidates.append(str(int(tier)))
    elif schema_type == "option" or "select" in custom:
        candidates.append({"value": str(int(tier))})
    else:
        candidates.extend([{"value": str(int(tier))}, str(int(tier)), int(tier)])

    deduped = []
    seen = set()
    for candidate in candidates:
        marker = repr(candidate)
        if marker not in seen:
            seen.add(marker)
            deduped.append(candidate)
    return deduped


def update_partner_tier(base_url, api_version, auth_type, username, token, issue_key, field_id, tier):
    editmeta = None
    try:
        editmeta = jira_get_issue_editmeta(base_url, api_version, auth_type, username, token, issue_key)
    except Exception:
        editmeta = None

    errors = []
    for payload in candidate_tier_payloads(editmeta, field_id, tier):
        try:
            jira_update_issue_fields(
                base_url, api_version, auth_type, username, token,
                issue_key, {field_id: payload},
            )
            return
        except Exception as error:
            errors.append(summarize_exception(error))

    raise RuntimeError(errors[-1] if errors else "Jira rejected the Partner tier update")


def _clear_partner_tier_state():
    for key in [
        "partner_tier_actionable", "partner_tier_blocked", "partner_tier_source_title",
        "partner_tier_source_rows", "partner_tier_field_ids", "partner_tier_apply_results",
        "partner_tier_confirm",
    ]:
        st.session_state.pop(key, None)


def render_partner_tier_maintenance():
    st.markdown(
        """<div class="feature-panel">
        <div class="step-kicker">Data maintenance · Partner master data</div>
        <div class="feature-title">Partner Tier Sync</div>
        <div class="feature-text">Export the current Partner Information page to Word, upload it here, review every proposed Jira change, and apply only the rows that are safe and unambiguous.</div>
        </div>""",
        unsafe_allow_html=True,
    )

    jira_base_url, api_version, auth_type, username, token, _ = current_jira_context()
    configured_url = str(PARTNER_TIER_CONFIG.get("confluence_page_url", "") or "")
    default_jql = str(PARTNER_TIER_CONFIG.get("jql", "") or DEFAULT_PARTNER_TIER_JQL)

    step1, step2, step3 = st.columns(3)
    with step1:
        st.markdown("""<div class="mini-card"><div class="step-kicker">Step 01</div><div class="tool-card-title">Export source</div><div class="tool-card-text">Confluence → Export to Word. Use the fresh file each week.</div></div>""", unsafe_allow_html=True)
    with step2:
        st.markdown("""<div class="mini-card"><div class="step-kicker">Step 02</div><div class="tool-card-title">Preview changes</div><div class="tool-card-text">The helper matches Partner / Project and extracts the group number safely.</div></div>""", unsafe_allow_html=True)
    with step3:
        st.markdown("""<div class="mini-card"><div class="step-kicker">Step 03</div><div class="tool-card-title">Confirm update</div><div class="tool-card-text">Blocked or ambiguous rows stay untouched. Jira updates require explicit confirmation.</div></div>""", unsafe_allow_html=True)

    st.write("")
    source_mode = "Upload file"
    uploaded_file = None
    page_url = ""

    source_col, action_col = st.columns([4.2, 1.1])
    with source_col:
        st.markdown("### Source file")
        st.caption("Weekly default: upload the newest Word export (.doc or .docx). Excel and CSV remain supported.")
        uploaded_file = st.file_uploader(
            "Partner Information file",
            type=["doc", "docx", "xlsx", "xlsm", "csv"],
            key="partner_tier_source_file",
            help="In Confluence choose Export to Word, then upload the exported .doc or .docx file.",
            label_visibility="collapsed",
        )
    with action_col:
        st.write("")
        st.write("")
        clear_clicked = st.button("Clear preview", use_container_width=True, key="partner_tier_clear")

    with st.expander("Advanced · direct Confluence source", expanded=False):
        st.caption("Use this only when Change Helper is running inside the corporate network. Streamlit Cloud cannot resolve the internal Confluence host.")
        use_direct_confluence = st.checkbox("Use Confluence URL instead of uploaded file", key="partner_tier_use_confluence")
        if use_direct_confluence:
            source_mode = "Confluence URL"
            page_url = st.text_input(
                "Confluence page URL",
                value=configured_url,
                key="partner_tier_confluence_url",
                placeholder="https://confluence.../display/...",
            )

    if uploaded_file is not None:
        st.success(f"Source selected: {uploaded_file.name}")

    with st.expander("JQL used to find empty Partner tier tickets", expanded=False):
        jql = st.text_area("JQL", value=default_jql, height=260, key="partner_tier_jql", label_visibility="collapsed")

    c1, c2, c3 = st.columns([1.2, 1.2, 3])
    preview_clicked = c1.button("Preview sync", type="primary", use_container_width=True, key="partner_tier_preview")
    c2.caption("Preview never changes Jira.")
    c3.caption("Missing, ambiguous, or malformed source rows are blocked instead of guessed.")

    if clear_clicked:
        _clear_partner_tier_state()
        st.rerun()

    if preview_clicked:
        _clear_partner_tier_state()
        if not require_jira_settings():
            return
        if source_mode == "Upload file" and uploaded_file is None:
            st.error("Upload the Partner Information file first. For Confluence, use Export to Word and upload the .doc or .docx.")
            return
        if source_mode == "Confluence URL" and not page_url.strip():
            st.error("Confluence page URL is empty.")
            return
        if not jql.strip():
            st.error("JQL is empty.")
            return

        partner_field_name = str(PARTNER_TIER_CONFIG.get("partner_field_name", "Partner / Project") or "Partner / Project")
        tier_field_name = str(PARTNER_TIER_CONFIG.get("partner_tier_field_name", "Partner tier") or "Partner tier")

        try:
            with st.spinner("Reading the source and Jira, then building a safe preview..."):
                if source_mode == "Upload file":
                    rows, source_title = load_partner_rows_from_upload(uploaded_file)
                else:
                    rows, source_title = load_confluence_partner_rows(
                        page_url,
                        jira_auth=(auth_type, username, token),
                    )
                lookup = build_partner_lookup(rows)

                field_ids = jira_get_field_ids_by_name(
                    jira_base_url, api_version, auth_type, username, token,
                    [partner_field_name, tier_field_name],
                )
                partner_field_id = field_ids.get(partner_field_name)
                tier_field_id = field_ids.get(tier_field_name)
                if not partner_field_id:
                    raise RuntimeError(f'Jira field "{partner_field_name}" was not found.')
                if not tier_field_id:
                    raise RuntimeError(f'Jira field "{tier_field_name}" was not found.')

                issues = jira_search(
                    jira_base_url, api_version, auth_type, username, token,
                    jql,
                    ["summary", "status", "updated", partner_field_id, tier_field_id],
                )
                actionable, blocked = build_preview_rows(
                    issues, jira_base_url, partner_field_id, tier_field_id, lookup
                )

                st.session_state.partner_tier_actionable = actionable
                st.session_state.partner_tier_blocked = blocked
                st.session_state.partner_tier_source_title = source_title
                st.session_state.partner_tier_source_rows = len(rows)
                st.session_state.partner_tier_field_ids = {
                    "partner": partner_field_id,
                    "tier": tier_field_id,
                }
        except Exception as error:
            st.error("Preview failed")
            st.code(summarize_exception(error, limit=1200))
            return

    actionable = st.session_state.get("partner_tier_actionable")
    blocked = st.session_state.get("partner_tier_blocked")
    if actionable is None and blocked is None:
        st.info("Upload the current source file and load a preview. Nothing is changed in Jira until the final confirmation step.")
        return

    actionable = actionable or []
    blocked = blocked or []
    total = len(actionable) + len(blocked)
    source_rows = int(st.session_state.get("partner_tier_source_rows", 0))
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Jira tickets", total)
    m2.metric("Ready to update", len(actionable))
    m3.metric("Blocked", len(blocked))
    m4.metric("Source rows", source_rows)
    st.caption(f'Source: **{st.session_state.get("partner_tier_source_title", "Uploaded file")}**')

    if actionable:
        st.subheader("Ready to update")
        st.caption("Matches are checked against both Partner and Project name. Only one safe source row/value is eligible.")
        action_df = st.data_editor(
            pd.DataFrame(actionable),
            hide_index=True,
            use_container_width=True,
            key="partner_tier_editor",
            disabled=["Key", "Open", "Summary", "Status", "Partner", "Source Partner", "Source Project", "Partner Category", "New tier", "Match"],
            column_config={
                "Select": st.column_config.CheckboxColumn("Select", default=True),
                "Open": st.column_config.LinkColumn("Open", display_text="Open"),
                "Summary": st.column_config.TextColumn("Summary", width="large"),
                "New tier": st.column_config.NumberColumn("New tier", format="%d"),
            },
        )
    else:
        action_df = pd.DataFrame()
        st.warning("No safe automatic updates found.")

    if blocked:
        with st.expander(f"Blocked / needs manual check ({len(blocked)})", expanded=True):
            st.dataframe(
                pd.DataFrame(blocked),
                hide_index=True,
                use_container_width=True,
                column_config={
                    "Open": st.column_config.LinkColumn("Open", display_text="Open"),
                    "Summary": st.column_config.TextColumn("Summary", width="large"),
                    "Reason": st.column_config.TextColumn("Reason", width="large"),
                },
            )

    selected_rows = []
    if not action_df.empty:
        selected_rows = action_df.loc[action_df["Select"] == True].to_dict("records")

    st.divider()
    confirmed = st.checkbox(
        f"I checked the preview and want to update {len(selected_rows)} selected ticket(s)",
        key="partner_tier_confirm",
    )
    apply_clicked = st.button(
        f"Update Partner tier ({len(selected_rows)})",
        type="primary",
        use_container_width=True,
        disabled=not confirmed or not selected_rows,
        key="partner_tier_apply",
    )

    if apply_clicked:
        tier_field_id = (st.session_state.get("partner_tier_field_ids") or {}).get("tier")
        if not tier_field_id:
            st.error("Partner tier field id is missing. Reload the preview.")
            return

        results = []
        progress = st.progress(0)
        status_box = st.empty()
        for index, row in enumerate(selected_rows, start=1):
            key = str(row.get("Key", "")).strip()
            tier = int(row.get("New tier"))
            status_box.info(f"Updating {index}/{len(selected_rows)} · {key} → tier {tier}")
            try:
                update_partner_tier(
                    jira_base_url, api_version, auth_type, username, token,
                    key, tier_field_id, tier,
                )
                results.append({"Key": key, "Tier": tier, "Status": "Success", "Details": "Partner tier updated"})
            except Exception as error:
                results.append({"Key": key, "Tier": tier, "Status": "Failed", "Details": summarize_exception(error)})
            progress.progress(index / len(selected_rows))

        st.session_state.partner_tier_apply_results = results
        success_count = sum(item["Status"] == "Success" for item in results)
        failed_count = len(results) - success_count
        if failed_count:
            status_box.warning(f"Finished: {success_count} updated, {failed_count} failed.")
        else:
            status_box.success(f"Finished: {success_count} ticket(s) updated successfully.")

    results = st.session_state.get("partner_tier_apply_results")
    if results:
        st.subheader("Update results")
        st.dataframe(pd.DataFrame(results), hide_index=True, use_container_width=True)

