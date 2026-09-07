import sys
import types
from io import BytesIO

try:
    import streamlit  # noqa: F401
except ModuleNotFoundError:
    fake = types.ModuleType("streamlit")
    fake.secrets = {}
    sys.modules["streamlit"] = fake

from docx import Document

from tools.partner_type_sync import (
    build_partner_type_lookup,
    match_partner_type,
    parse_partner_type_rows_from_doc,
    parse_partner_type_rows_from_docx,
    _rows_from_matrix,
)


def test_partner_type_word_parser():
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    table.cell(0, 0).text = "Partner"
    table.cell(0, 1).text = "Project name"
    table.cell(0, 2).text = "Partner Category"
    table.cell(0, 3).text = "Solution type"
    table.cell(1, 0).text = "GrandPashaBet"
    table.cell(1, 1).text = "Betsin"
    table.cell(1, 3).text = "Builder (Turnkey)"
    table.cell(2, 0).text = "GrandPashaBet"
    table.cell(2, 1).text = "Solibet"
    table.cell(2, 3).text = "Casweb (Turnkey)"

    buffer = BytesIO()
    doc.save(buffer)
    rows = parse_partner_type_rows_from_docx(buffer.getvalue())

    assert rows[0] == {
        "Partner": "GrandPashaBet",
        "Project name": "Betsin",
        "Partner Type": "Builder (Turnkey)",
    }
    assert rows[1]["Project name"] == "Solibet"
    assert rows[1]["Partner Type"] == "Casweb (Turnkey)"


def test_partner_type_matching_tolerates_spacing():
    lookup = build_partner_type_lookup([
        {"Partner": "1 Bet Pro", "Project name": "", "Partner Type": "B2B"},
    ])
    entry, error, match = match_partner_type(lookup, "1BetPro")
    assert error == ""
    assert entry["partner_type"] == "B2B"
    assert match.startswith("Partner normalized")


def test_partner_type_matches_project_name():
    lookup = build_partner_type_lookup([
        {"Partner": "GrandPashaBet", "Project name": "Betsin", "Partner Type": "Builder (Turnkey)"},
        {"Partner": "GrandPashaBet", "Project name": "Solibet", "Partner Type": "Casweb (Turnkey)"},
    ])
    entry, error, match = match_partner_type(lookup, "Solibet")
    assert error == ""
    assert entry["project"] == "Solibet"
    assert entry["partner_type"] == "Casweb (Turnkey)"
    assert match.startswith("Project exact")


def test_partner_type_prefixed_project_suffix():
    lookup = build_partner_type_lookup([
        {"Partner": "GrandPashaBet", "Project name": "Solibet", "Partner Type": "Casweb (Turnkey)"},
    ])
    entry, error, match = match_partner_type(lookup, "albatross-Solibet")
    assert error == ""
    assert entry["project"] == "Solibet"
    assert "Suffix" in match


def test_partner_type_all_prefers_same_named_project():
    lookup = build_partner_type_lookup([
        {"Partner": "GrandPashaBet", "Project name": "Betsin", "Partner Type": "Builder (Turnkey)"},
        {"Partner": "GrandPashaBet", "Project name": "GrandPashaBet", "Partner Type": "Casweb (Turnkey)"},
        {"Partner": "GrandPashaBet", "Project name": "Solibet", "Partner Type": "Casweb (Turnkey)"},
    ])
    entry, error, match = match_partner_type(lookup, "GrandPashaBet-ALL")
    assert error == ""
    assert entry["project"] == "GrandPashaBet"
    assert entry["partner_type"] == "Casweb (Turnkey)"
    assert "ALL" in match


def test_partner_type_all_without_same_named_project_blocks_mixed_types():
    lookup = build_partner_type_lookup([
        {"Partner": "Umbrella", "Project name": "One", "Partner Type": "Builder"},
        {"Partner": "Umbrella", "Project name": "Two", "Partner Type": "Casweb"},
    ])
    entry, error, _ = match_partner_type(lookup, "Umbrella-ALL")
    assert entry is None
    assert "Conflicting" in error


def test_empty_partner_type_is_blocked():
    lookup = build_partner_type_lookup([
        {"Partner": "188bet", "Project name": "", "Partner Type": ""},
    ])
    entry, error, _ = match_partner_type(lookup, "188bet")
    assert entry is None
    assert "empty" in error.lower()


def test_conflicting_partner_types_are_blocked():
    lookup = build_partner_type_lookup([
        {"Partner": "Same Partner", "Project name": "A", "Partner Type": "B2B"},
        {"Partner": "Same Partner", "Project name": "B", "Partner Type": "White Label"},
    ])
    entry, error, _ = match_partner_type(lookup, "Same Partner")
    assert entry is None
    assert "Conflicting" in error


def test_duplicate_project_across_parents_is_blocked():
    lookup = build_partner_type_lookup([
        {"Partner": "Parent A", "Project name": "Shared", "Partner Type": "Casweb"},
        {"Partner": "Parent B", "Project name": "Shared", "Partner Type": "Casweb"},
    ])
    entry, error, _ = match_partner_type(lookup, "Shared")
    assert entry is None
    assert "multiple partners" in error


def test_partner_type_confluence_doc_html_export_with_rowspan():
    data = b"""<html><body><table>
    <tr><th>Partner</th><th>Project name</th><th>Partner Category</th><th>Solution type</th></tr>
    <tr><td rowspan=\"2\">GrandPashaBet</td><td>Betsin</td><td>Group 4</td><td>Builder (Turnkey)</td></tr>
    <tr><td>Solibet</td><td>Group 4</td><td>Casweb (Turnkey)</td></tr>
    </table></body></html>"""
    rows = parse_partner_type_rows_from_doc(data)
    assert rows[0]["Partner"] == "GrandPashaBet"
    assert rows[0]["Project name"] == "Betsin"
    assert rows[1]["Partner"] == "GrandPashaBet"
    assert rows[1]["Project name"] == "Solibet"
    assert rows[1]["Partner Type"] == "Casweb (Turnkey)"


def test_old_partner_type_header_is_still_accepted():
    matrix = [
        ["Partner", "Project name", "Partner Type"],
        ["LegacyPartner", "", "B2B"],
    ]
    rows = _rows_from_matrix(matrix)
    assert rows == [{"Partner": "LegacyPartner", "Project name": "", "Partner Type": "B2B"}]


def test_partner_type_repeated_partner_project_pair_collapses_to_partner():
    lookup = build_partner_type_lookup([
        {"Partner": "Festwin", "Project name": "-", "Partner Type": "Casweb (Turnkey)"},
    ])
    entry, error, match = match_partner_type(lookup, "Festwin - Festwin")
    assert error == ""
    assert entry["partner"] == "Festwin"
    assert entry["partner_type"] == "Casweb (Turnkey)"
    assert "Repeated pair" in match


def test_partner_type_repeated_pair_with_unicode_dash_collapses_to_partner():
    lookup = build_partner_type_lookup([
        {"Partner": "Festwin", "Project name": "-", "Partner Type": "Casweb (Turnkey)"},
    ])
    entry, error, match = match_partner_type(lookup, "Festwin – Festwin")
    assert error == ""
    assert entry["partner_type"] == "Casweb (Turnkey)"
    assert "Repeated pair" in match
