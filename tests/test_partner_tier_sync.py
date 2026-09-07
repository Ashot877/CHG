import sys
import types
from io import BytesIO

from docx import Document

try:
    import streamlit  # noqa: F401
except ModuleNotFoundError:
    fake = types.ModuleType("streamlit")
    fake.secrets = {}
    sys.modules["streamlit"] = fake

from tools.partner_tier_sync import (
    build_partner_lookup,
    extract_tier_from_category,
    match_partner,
    normalize_partner,
    parse_partner_rows_from_html,
    parse_partner_rows_from_docx,
)


def test_category_variations():
    cases = {
        "Group 4": 4,
        "Group 4(New)": 4,
        " group   4 ": 4,
        "GROUP 4": 4,
        "Tier 4": 4,
    }
    for raw, expected in cases.items():
        tier, error = extract_tier_from_category(raw)
        assert tier == expected
        assert error == ""


def test_empty_and_ambiguous_category_are_blocked():
    assert extract_tier_from_category("")[0] is None
    assert extract_tier_from_category("Group A")[0] is None
    assert extract_tier_from_category("4 / 5")[0] is None


def test_partner_matching_tolerates_spacing_but_not_ambiguity():
    lookup = build_partner_lookup([
        {"Partner": "1 Bet Pro", "Partner Category": "Group 5"},
        {"Partner": "188bet", "Partner Category": "Group 4(New)"},
    ])
    entry, error, match = match_partner(lookup, "1BetPro")
    assert error == ""
    assert entry["tier"] == 5
    assert match == "Normalized"


def test_conflicting_duplicate_is_blocked():
    lookup = build_partner_lookup([
        {"Partner": "SamePartner", "Partner Category": "Group 3"},
        {"Partner": "SamePartner", "Partner Category": "Group 4"},
    ])
    entry, error, _ = match_partner(lookup, "SamePartner")
    assert entry is None
    assert "Conflicting" in error


def test_html_table_parser():
    html = """
    <table>
      <tr><th>Partner</th><th>Project name</th><th>Partner Category</th></tr>
      <tr><td>1betpro</td><td>-</td><td> Group 5 </td></tr>
      <tr><td>188bet</td><td>-</td><td>Group 4(New)</td></tr>
    </table>
    """
    rows = parse_partner_rows_from_html(html)
    assert rows[0]["Partner"] == "1betpro"
    assert rows[1]["Partner Category"] == "Group 4(New)"


def test_word_export_table_parser():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Partner"
    table.cell(0, 1).text = "Project name"
    table.cell(0, 2).text = "Partner Category"
    table.cell(1, 0).text = "1betpro"
    table.cell(1, 1).text = "-"
    table.cell(1, 2).text = "Group 5"
    table.cell(2, 0).text = "188bet"
    table.cell(2, 1).text = "-"
    table.cell(2, 2).text = "Group 4(New)"

    buffer = BytesIO()
    doc.save(buffer)
    rows = parse_partner_rows_from_docx(buffer.getvalue())
    assert rows[0] == {"Partner": "1betpro", "Partner Category": "Group 5"}
    assert rows[1]["Partner Category"] == "Group 4(New)"
